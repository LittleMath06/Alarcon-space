import csv
import io

import pytest
from openpyxl import load_workbook

from espaco_alarcon.db import conectar, migrar
from espaco_alarcon.domain import leads, painel
from espaco_alarcon.domain.busca import Filtros, buscar
from espaco_alarcon.domain.exportar import exportar
from espaco_alarcon.jobs.importar_receita import Fonte, Importador

from conftest import criar_usuario, csrf, fazer_login, novo_cliente
from receita_ficticia import gerar


@pytest.fixture
def base(cfg, tmp_path):
    conn = conectar(cfg.db_path)
    migrar(conn)
    Importador(cfg, conn, Fonte(str(gerar(tmp_path / "zips", extras_estab=10)), "local"), competencia="2026-09").executar()
    yield conn
    conn.close()


def test_pipeline_completo(base):
    lid = leads.salvar(base, "11111111000100")
    assert leads.salvar(base, "11111111000100") == lid             # sem duplicar
    assert base.execute("SELECT COUNT(*) FROM lead").fetchone()[0] == 1
    assert leads.obter(base, lid)["etapa_nome"] == "Novo"
    leads.mover(base, lid, 3)
    leads.favoritar(base, lid, True)
    leads.lembrete(base, lid, "2020-01-01")
    leads.anotar(base, lid, "Pediu catálogo de cabo automotivo.")
    with pytest.raises(ValueError):
        leads.anotar(base, lid, "   ")
    with pytest.raises(ValueError):
        leads.lembrete(base, lid, "32/13/2026")
    with pytest.raises(ValueError):
        leads.mover(base, lid, 99)
    assert [l["id"] for l in leads.listar(base, filtro="lembretes")] == [lid]
    assert [l["id"] for l in leads.listar(base, filtro="favoritos")] == [lid]
    with pytest.raises(ValueError):
        leads.descartar(base, lid, "")
    leads.descartar(base, lid, "Já compra de concorrente")
    with pytest.raises(ValueError):
        leads.mover(base, lid, 5)
    assert leads.listar(base) == [] and len(leads.listar(base, filtro="descartados")) == 1
    leads.reabrir(base, lid)
    leads.mover(base, lid, 5)                                         # cliente só por ação do usuário
    ficha = leads.do_estabelecimento(base, "11111111000100")
    assert ficha["cliente"] == 1 and len(ficha["notas"]) == 1
    eventos = [h["evento"] for h in ficha["historico"]]
    assert {"criado", "etapa", "favorito", "lembrete", "descartado", "reaberto"} <= set(eventos)
    with pytest.raises(LookupError):
        leads.salvar(base, "00000000000000")


def test_busca_por_situacao_de_lead(base):
    total = buscar(base, Filtros(uf="SP")).total
    lid = leads.salvar(base, "12121212000100")
    assert buscar(base, Filtros(uf="SP", lead_status="salvo")).total == 1
    assert buscar(base, Filtros(uf="SP", lead_status="nao_salvo")).total == total - 1
    leads.descartar(base, lid, "fora do perfil")
    assert buscar(base, Filtros(uf="SP")).total == total - 1             # descartado some por padrão
    assert buscar(base, Filtros(uf="SP", lead_status="descartado")).total == 1


def test_exportacao_csv_xlsx_e_formula(tmp_path):
    linhas = [{"a": "=HYPERLINK(\"x\")", "b": "Refrigeração Ômega"}, {"a": "normal", "b": None}]
    colunas = [("a", "Coluna A"), ("b", "Coluna B")]
    c = exportar(linhas, colunas, tmp_path, "teste", "csv")
    texto = c.read_bytes()
    assert texto.startswith(b"\xef\xbb\xbf")                          # BOM para o Excel
    linhas_csv = list(csv.reader(io.StringIO(texto.decode("utf-8-sig")), delimiter=";"))
    assert linhas_csv[0] == ["Coluna A", "Coluna B"] and linhas_csv[1][0].startswith("'=")
    x = exportar(linhas, colunas, tmp_path, "teste", "xlsx")
    ws = load_workbook(x).active
    assert ws["A2"].value.startswith("'=") and ws["B2"].value == "Refrigeração Ômega"
    with pytest.raises(ValueError):
        exportar(linhas, colunas, tmp_path, "teste", "pdf")


def test_metricas_reais(base):
    m = painel.metricas(base)
    assert m["estabelecimentos"] == base.execute("SELECT COUNT(*) FROM estabelecimento").fetchone()[0]
    assert m["leads_salvos"] == 0 and m["clientes"] == 0
    leads.mover(base, leads.salvar(base, "11111111000100"), 5)
    m = painel.metricas(base)
    assert m["leads_salvos"] == 1 and m["clientes"] == 1
    assert sum(e["n"] for e in m["por_etapa"]) == 1
    assert sum(r["n"] for r in m["por_regiao"]) == m["estabelecimentos"]
    assert all(t["cnpj"] != "11111111000100" for t in m["top"])           # já trabalhado não aparece no top
    assert {f["fonte"] for f in painel.status_fontes(base)} >= {"Receita Federal (Dados Abertos do CNPJ)"}


def test_telas_leads_exportacao_e_dashboard(cfg, base):
    c = novo_cliente(cfg)
    criar_usuario(c)
    fazer_login(c)
    r = c.get("/")
    assert "leads novos" in r.text and "estabelecimentos na base" in r.text
    token = csrf(r.text)
    r = c.get("/busca", params={"uf": "SP", "segmento": "1"})
    assert "Salvar lead" in r.text
    r = c.post("/leads/salvar", data={"csrf_token": token, "cnpj": "11111111000100", "voltar": "/busca?uf=SP"})
    assert r.headers["location"] == "/busca?uf=SP"
    r = c.post("/leads/salvar", data={"csrf_token": token, "cnpj": "11111111000211", "voltar": "https://malicioso.invalid/"})
    assert r.headers["location"] == "/leads"                           # sem redirecionamento aberto
    assert c.post("/leads/salvar", data={"cnpj": "11111111000100"}).status_code == 403
    r = c.get("/leads")
    assert "EMPRESA FICTICIA 11111111 LTDA" in r.text and "Novo" in r.text
    lid = base.execute("SELECT id FROM lead WHERE cnpj = '11111111000100'").fetchone()[0]
    r = c.post(f"/leads/{lid}/descartar", data={"csrf_token": token, "motivo": ""}, headers={"referer": "http://testserver/leads"})
    assert "lead_erro" in r.headers["location"]
    assert c.post("/leads/999/etapa", data={"csrf_token": token, "etapa_id": "2"}).status_code == 404
    r = c.get("/leads/exportar", params={"formato": "xlsx"})
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert list((cfg.data_dir / "exports").glob("leads-*.xlsx"))
    r = c.get("/busca/exportar", params={"uf": "SP", "formato": "csv"})
    texto = r.content.decode("utf-8-sig")
    assert r.status_code == 200 and texto.splitlines()[0].startswith("CNPJ;Razão social")
    assert "SOCIO" not in texto                                         # sem dados de sócios
    assert c.get("/busca/exportar", params={"formato": "csv"}).status_code == 400
    assert c.get("/busca/exportar", params={"uf": "SP", "formato": "exe"}).status_code == 400
    assert "Salvar como lead" in c.get("/empresa/12121212000100").text
    assert "Lead" in c.get("/empresa/11111111000100").text
